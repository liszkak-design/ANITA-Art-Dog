#!/usr/bin/env python3
"""Partie optymalizacji płytek ceramicznych "with a Photo" (299 ofert z plans/mapa_koncowa.csv, decyzja OPTYMALIZUJ).
  python3 -I scripts/batch.py ids                 # pobiera ID i tytuły ofert sklepu -> plans/listing_ids_sklep.csv (odczyt)
  python3 -I scripts/batch.py plan N              # dry-run partii N (max 50): plans/partia_00N.csv, _dryrun.json, partie_status.csv
  python3 -I scripts/batch.py apply N --apply     # zapis do Etsy; tylko po wyraźnym "tak" dla partii N, token z listings_w
Każda oferta jest zapisywana raz: oferta z wpisem "weryfikacja" ok w jakimkolwiek logs/zapis_<id>_*.jsonl jest pomijana.
Pominięte (decyzja 26.10.2026): SKU, shopping basket summary, GPSR (ustawia właścicielka ręcznie w Etsy).
"""
import csv
import datetime
import glob
import html
import json
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import etsy_api as api
import generators as gen
import translations as T
import apply_test_sphynx as base

ROOT = api.ROOT
SHOP_ID = base.SHOP_ID
BATCH_MAX = 50
TODAY = datetime.date.today()
IDS_CSV = ROOT / "plans" / "listing_ids_sklep.csv"
APPROVED_CSV = ROOT / "plans" / "zatwierdzone_dopasowania_zdjec.csv"  # oferty dopasowane przez zdjęcia, zatwierdzone przez właścicielkę
ID_IMG_CSV = ROOT / "plans" / "id_po_obrazie.csv"  # cache: oferta z eksportu -> ID w Etsy rozstrzygnięte po zdjęciu
STATUS_CSV = ROOT / "plans" / "partie_status.csv"
EAN_BASE = "https://www.artdog.com.pl/allegro/zebraneEAN/"
CACHE = base.CACHE.parent / "batch"
IMAGE_ORDER = base.IMAGE_ORDER
LANGS = base.LANGS
MATERIAL_PROPERTY, CERAMIC_VALUE = 148789511893, 83
UA = {"User-Agent": "ArtDog-Etsy-Tool/1.0"}  # serwer artdog.com.pl odrzuca domyślny UA Pythona (403)
FEAT = re.compile(r"long|short|wire|smooth|fohlen|cropped|curly", re.I)
OVERRIDES_CSV = ROOT / "plans" / "dopasowania_wlasciciela.csv"  # rozstrzygnięcia właścicielki: oferta (ID) -> rasa, odmiana, EAN
EXCLUDE_BREEDS = {"Pointer", "German Pointer"}
SP_EN = {"kot": "cat", "pies": "dog", "koń": "horse", "kucyk": "pony"}
PL_FORMS = {"kot": dict(gen="kota", loc="kocie"), "pies": dict(gen="psa", loc="psie"), "koń": dict(gen="konia", loc="koniu"),
            "kucyk": dict(gen="kucyka", loc="kucyku")}


# ---------- dane wejściowe ----------
def rasy():
    ws = openpyxl.load_workbook(ROOT / "data" / "Rasy.xlsx", data_only=True)["Rasy"]
    h = [c.value for c in ws[1]]
    return {(r[1], r[0]): dict(zip(h, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}


def owner_overrides():
    return list(csv.DictReader(open(OVERRIDES_CSV, encoding="utf-8"))) if OVERRIDES_CSV.exists() else []


def species_of(row):
    return "kucyk" if row["gatunek"] == "koń" and "pony" in row["rasa_kolumna_C"].lower() else row["gatunek"]


def feature_words():
    """Słowa odmian z arkusza 'Słowa stałe': klucz EN -> {kod języka: słowo}. Nawiasy z uwagami usuwamy, poza niemieckim małe litery."""
    ws = openpyxl.load_workbook(ROOT / "data" / "Rasy.xlsx", data_only=True)["Słowa stałe"]
    cols = dict(nl=2, fr=4, de=6, it=8, pl=10, pt=12, es=14, sv=16)
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0]:
            out[r[0]] = {c: (lambda w: w if c == "de" else w.lower())(re.sub(r"\s*\(.*?\)", "", str(r[i] or "")).strip()) for c, i in cols.items()}
    # "foal" (źrebię) z arkusza "Opisy statuetek": kolumny PL, DE, FR, ES, NL, IT, PT, SV
    for r in openpyxl.load_workbook(ROOT / "data" / "Rasy.xlsx", data_only=True)["Opisy statuetek"].iter_rows(min_row=2, values_only=True):
        if r[0] and str(r[0]).startswith("foal"):
            out["foal"] = {c: (lambda w: w if c == "de" else w.lower())(str(r[i] or "").strip()) for c, i in dict(pl=5, de=6, fr=7, es=8, nl=9, it=10, pt=11, sv=12).items()}
    return out


def feature_en(key):
    """Angielskie słowo odmiany do tytułu: 'Fox Terrier Smooth' -> 'Smooth', 'Dachshund Wire-haired' -> 'Wire-haired', 'cropped ears' -> 'Cropped Ears'."""
    if key == "foal":
        return "Foal"
    return "Cropped Ears" if key.lower().endswith("ears") and key.lower().startswith("cropped") else key.split()[-1]


def feature_residue(row):
    """Tekst nazwy z oferty i folderu po odjęciu nazwy rasy z kolumny C (reszta to ewentualna odmiana: Longhaired, Fohlen, Cropped)."""
    folder = re.sub(r"^\d+\.?\s*|\s*\d{13}\s*$", "", row["folder"])
    return re.sub(re.escape(row["rasa_kolumna_C"]), "", row["nazwa_w_tytule"] + " " + folder, flags=re.I)


def e_class():
    """Klasyfikacja wpisów kolumny E (data/nick_klasyfikacja.csv): rasa_id -> 'zdrobnienie' lub 'synonim'."""
    return {r["rasa_id"]: r["typ"] for r in csv.DictReader(open(ROOT / "data" / "nick_klasyfikacja.csv", encoding="utf-8"))}


def approved_photo_matches():
    return {r["zrodlo"]: r for r in csv.DictReader(open(APPROVED_CSV, encoding="utf-8"))} if APPROVED_CSV.exists() else {}


def img_key(url):
    m = re.search(r"/il/([0-9a-f]+)/(\d+)/", url or "")
    return m.groups() if m else None


def export_image_url(zrodlo):
    f, n = zrodlo.split(":")
    return list(csv.DictReader(open(ROOT / "data" / f, encoding="utf-8-sig")))[int(n) - 2]["IMAGE1"]


def resolve_id_by_image(zrodlo, ids):
    """Oferty o identycznym tytule: ID wybieramy po zdjęciu (adres zdjęcia z eksportu = adres zdjęcia oferty w Etsy). Wynik w cache."""
    cache = {r["zrodlo"]: r for r in csv.DictReader(open(ID_IMG_CSV, encoding="utf-8"))} if ID_IMG_CSV.exists() else {}
    if zrodlo in cache:
        return int(cache[zrodlo]["listing_id"]) if cache[zrodlo]["listing_id"] else None
    want, hit = img_key(export_image_url(zrodlo)), []
    for lid in ids:
        imgs = api.get(f"/v3/application/listings/{lid}/images").get("results", [])
        if any(img_key(i.get("url_fullxfull")) == want for i in imgs):
            hit.append(lid)
    res = hit[0] if len(hit) == 1 else None
    cache[zrodlo] = {"zrodlo": zrodlo, "listing_id": res or "", "kandydaci": "/".join(map(str, ids))}
    with open(ID_IMG_CSV, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=["zrodlo", "listing_id", "kandydaci"])
        w.writeheader()
        w.writerows(cache.values())
    return res


def read_map():
    rows = [r for r in csv.DictReader(open(ROOT / "plans" / "mapa_koncowa.csv", encoding="utf-8")) if r["decyzja"] == "OPTYMALIZUJ"]
    ean = {r["zrodlo"]: r["ean"] for r in csv.DictReader(open(ROOT / "plans" / "mapa_plytki_with_photo.csv", encoding="utf-8"))}
    appr = approved_photo_matches()
    for r in rows:
        r["ean"] = ean.get(r["zrodlo"], "")
        if not r["ean"] and r["zrodlo"] in appr:  # dopasowanie po zdjęciach zatwierdzone przez właścicielkę: EAN z nazwy folderu
            r["ean"] = appr[r["zrodlo"]]["ean_z_folderu"]
    return rows


def fetch_ids():
    out, off = [], 0
    while True:
        r = api.get(f"/v3/application/shops/{SHOP_ID}/listings?state=active&limit=100&offset={off}")
        if "results" not in r:
            sys.exit("BŁĄD pobierania ofert: " + str(r)[:200])
        out += [(x["listing_id"], html.unescape(x["title"]), x["state"]) for x in r["results"]]
        off += 100
        if off >= r["count"]:
            break
    with open(IDS_CSV, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["listing_id", "title", "state"])
        w.writerows(out)
    print("zapisano", len(out), "ofert w", IDS_CSV.relative_to(ROOT))


def title_index():
    idx = {}
    for r in csv.DictReader(open(IDS_CSV, encoding="utf-8")):
        idx.setdefault(r["title"], []).append(int(r["listing_id"]))
    return idx


def sales():
    cnt = {}
    rd = csv.DictReader(open(ROOT / "data" / "EtsySoldOrderItems2025-2026.csv", encoding="cp1252"), delimiter=";")
    for r in rd:
        try:
            cnt[int(r["Listing ID"])] = cnt.get(int(r["Listing ID"]), 0) + int(r["Quantity"] or 1)
        except ValueError:
            pass
    return cnt


def is_done(lid):
    for f in glob.glob(str(ROOT / "logs" / f"zapis_{lid}_*.jsonl")):
        for line in open(f):
            j = json.loads(line)
            if j["id"] == "weryfikacja" and j["status"] == "ok":
                return True
    return False


# ---------- treści ----------
ALT_TEMPLATES = None


def alt_templates():
    global ALT_TEMPLATES
    if ALT_TEMPLATES is None:
        en = base.parse_en()
        ALT_TEMPLATES = {}
        for f, alt in en["alts"].items():
            a = alt.replace("a Sphynx cat", "{A_X}").replace("Sphynx cat", "{X}")
            assert "Sphynx" not in a, a
            ALT_TEMPLATES[f] = a
    return ALT_TEMPLATES


def article(word):
    return "an" if word[0].upper() in "AEIOU" and not word.lower().startswith(("uni", "eu", "use")) else "a"


def alts(breed, species):
    sp = SP_EN[species]
    x = breed if sp in breed.lower() else f"{breed} {sp}"
    return {f: t.replace("{A_X}", f"{article(x)} {x}").replace("{X}", x) for f, t in alt_templates().items()}


def parse_pl():
    t = (ROOT / "templates" / "tlumaczenie_pl_plytka.md").read_text(encoding="utf-8")
    titles = re.findall(r"^\d\. (\{B\}.*)$", base.section(t, "Tytuł (6 wzorców"), re.M)
    banks = {s: re.findall(r"^- " + s + r"\d: (.*)$", base.section(t, "Pierwszy akapit: banki"), re.M) for s in "ABC"}
    fixed = base.codeblock(base.section(t, "Sekcje stałe"))
    tg = base.section(t, "Tagi (max 20 znaków)")
    stal, pula = tg.split("**Pula**")
    core = re.findall(r"→ (.+?) \(\d+\)$", stal, re.M)
    spec = {}
    for k, w in (("kot", "cat"), ("pies", "dog"), ("koń", "horse"), ("kucyk", "pony")):
        m = re.search(r"^- " + w + r": (.+?) \(\d+\), (.+?) \(\d+\)$", stal, re.M)
        spec[k] = (m.group(1), m.group(2))
    return dict(titles=titles, A=banks["A"], B=banks["B"], C=banks["C"], fixed=fixed, core=core, spec=spec,
                pool=re.findall(r"→ (.+?) \(\d+\)$", pula, re.M))


PL = None


def render_translation(code, lid, species, row, ch, en_tags, title_idx, feat=""):
    """Tłumaczenie z szablonów i słownika; indeksy wariantów jak w angielskim (N-warianty zdrobnień: losowo wśród 4 zwykłych)."""
    global PL
    pick = lambda slot, n: ch[slot] if ch[slot] < n else gen.pick(lid, slot + "t", n)
    if code == "pl":
        PL = PL or parse_pl()
        d, b = PL, row["PL nazwa"]
        sp = PL_FORMS[species]
        fill = lambda t, bb=None: re.sub(r"\s+", " ", t.replace("{B}", bb or b).format(**sp)).strip()
        fixed = d["fixed"].replace("{B}", b)
        for k, v in sp.items():
            fixed = fixed.replace("{" + k + "}", v)
    else:
        d = T.L[code]
        b = row[d["col"]]
        sp = d["sp"][species]
        fill = lambda t, bb=None: T.fill(t, sp, bb or b)
        fixed = T.fill_block(d["fixed"], sp, b)
    par = " ".join(fill(d[s][pick(s, len(d[s]))]) for s in "ABC")
    tg = list(d["core"]) + list(d["spec"][species]) + [d["pool"][T.POOL_EN.index(x)] for x in en_tags[7:]]
    title_name = (b + " " + feat).strip() if feat else None  # odmiana tylko w tytule, nie w opisie
    return {"title": fill(d["titles"][title_idx - 1], title_name), "description": par + "\n\n" + fixed, "tags": tg}


def render(row, lid, R):
    g0, sp = row["gatunek"], species_of(row)
    breed = row["rasa_kolumna_C"]
    rr = R[(g0, row["rasa_id"])]
    fkey = row.get("klucz_odmiany", "")
    feat_en = feature_en(fkey) if fkey else ""
    e_val, typ = (rr.get("Potoczne EN") or "").strip(), e_class().get(row["rasa_id"], "")
    nick, syn = (e_val if typ == "zdrobnienie" else ""), (e_val if typ == "synonim" else "")
    desc, ch = gen.description(lid, sp, breed, nick, syn)
    # banki mają na sztywno "a {BREED}" / "a {SYN}": poprawna forma a/an zależnie od pierwszej litery nazwy
    for name in {n for n in (breed, syn) if n}:
        desc = re.sub(r"\b([Aa]) (?=" + re.escape(name) + r")",
                      lambda m, name=name: (article(name).capitalize() if m.group(1) == "A" else article(name)) + " ", desc)
    # "Hairless Dog dog", "Sheepdog dog", "Ocicat cat" -> bez powtórzenia
    desc = re.sub(r"\b(\w*(?:dog|cat|horse|pony))\s+(?:dog|cat|horse|pony)\b",
                  lambda m: m.group(1) if m.group(1).lower().endswith(m.group(0).split()[-1].lower()) else m.group(0), desc, flags=re.I)
    if syn:
        assert syn in desc.split("\n\n")[0], f"brak synonimu w pierwszym akapicie: {lid}"
    assert breed in desc, f"brak pełnej nazwy rasy w opisie: {lid}"  # zdanie A może być samym zdrobnieniem, rasa jest w IN SHORT
    tit, tidx = gen.title(lid, sp, breed, feat_en)
    tg = gen.tags(lid, sp)
    out = {"title": tit, "description": desc, "tags": tg, "materials": ["ceramic"],
           "alts": alts((breed + " " + feat_en).strip(), sp), "wzorzec_tytulu": tidx, "ABC": ch, "nick": nick, "synonim": syn,
           "odmiana": fkey}
    fw = feature_words().get(fkey, {}) if fkey else {}
    if fkey and not fw:
        raise ValueError(f"brak słów odmiany '{fkey}' w arkuszu Słowa stałe")
    trans = {c: render_translation(c, lid, sp, rr, ch, tg, tidx, fw.get(c, "")) for c in LANGS}
    return out, trans


def check(en, trans, ean):
    prob = []
    if len(en["title"]) > 140:
        prob.append(f"tytuł EN {len(en['title'])} zn.")
    if len(en["tags"]) != 13 or len(set(en["tags"])) != 13 or any(len(x) > 20 for x in en["tags"]):
        prob.append("tagi EN")
    if not ean:
        prob.append("brak EAN")
    for c, d in trans.items():
        if len(d["title"]) > 140:
            prob.append(f"tytuł {c} {len(d['title'])} zn.")
        long_ = [x for x in d["tags"] if len(x) > 20]
        if len(d["tags"]) != 13 or long_:
            prob.append(f"tagi {c}: {long_ or len(d['tags'])}")
        if "{" in d["title"] + d["description"] or "None" in d["title"]:
            prob.append(f"nierozwiązany znacznik w {c}")
    if "{" in en["title"] + en["description"]:
        prob.append("nierozwiązany znacznik EN")
    return prob


def reachable(ean):
    for f in IMAGE_ORDER[:-1]:
        err = ""
        for attempt in range(3):  # chwilowe błędy serwera zdjęć nie wykluczają oferty
            try:
                with urllib.request.urlopen(urllib.request.Request(f"{EAN_BASE}{ean}/{f}", method="HEAD", headers=UA), timeout=30) as r:
                    err = "" if r.status == 200 else f"{f}: {r.status}"
                break
            except Exception as e:
                err = f"{f}: {str(e)[:60]}"
                time.sleep(2 * (attempt + 1))
        if err:
            return err
    return ""


# ---------- plan ----------
def build_candidates():
    R = rasy()
    idx = title_index()
    sold = sales()
    cands = []
    for row in read_map():
        why = []
        ids = idx.get(row["tytul"], [])
        if len(ids) > 1 and row["zrodlo"] in approved_photo_matches():
            r_id = resolve_id_by_image(row["zrodlo"], ids)
            ids = [r_id] if r_id else ids
        if len(ids) != 1:
            why.append(f"ID oferty: {len(ids)} dopasowań po tytule")
        lid = ids[0] if len(ids) == 1 else 0
        if lid and is_done(lid):
            why.append("już zoptymalizowana")
        res = feature_residue(row)
        if FEAT.search(res):
            if re.fullmatch(r"\W*cropped\W*", res, re.I):
                row["klucz_odmiany"] = "cropped ears"  # Schnauzer Cropped: słowa w arkuszu Słowa stałe (decyzja właścicielki: odmiana w tytule, nie w opisie)
            else:
                why.append("odmiana sierści/wiek w nazwie: brak słów w arkuszu, wymaga decyzji")
        if row["do_decyzji"]:
            why.append("do decyzji: " + row["do_decyzji"])
        if row["rasa_kolumna_C"] in EXCLUDE_BREEDS:
            why.append("Pointer: sprawdza właścicielka")
        rr = R.get((row["gatunek"], row["rasa_id"]))
        if not rr:
            why.append("brak rasy w Rasy.xlsx")
        elif any(not rr.get(c) for c in ["PL nazwa"] + [T.L[l]["col"] for l in LANGS if l != "pl"]):
            why.append("brak tłumaczenia nazwy rasy")
        cands.append({"row": row, "lid": lid, "sales": sold.get(lid, 0), "why": why})
    # rozstrzygnięcia właścicielki (oferta wskazana ID): wchodzą zamiast wierszy mapy o tym samym ID lub z niejednoznacznym ID tej rasy
    ov = [o for o in owner_overrides() if o["aktywne"] == "tak"]
    if ov:
        ovl, covered = {int(o["listing_id"]) for o in ov}, {o["rasa_kolumna_C"] for o in ov}
        cands = [c for c in cands if c["lid"] not in ovl and not (c["lid"] == 0 and c["row"]["rasa_kolumna_C"] in covered)]
        byid = {int(r["listing_id"]): r["title"] for r in csv.DictReader(open(IDS_CSV, encoding="utf-8"))}
        by_name = {v["Nazwa EN (główna)"]: (k[0], k[1]) for k, v in R.items()}
        for o in ov:
            lid, why = int(o["listing_id"]), []
            g0, rid = by_name[o["rasa_kolumna_C"]]
            row = {"zrodlo": f"wlasciciel:{lid}", "gatunek": g0, "nazwa_w_tytule": o["rasa_kolumna_C"], "rasa_kolumna_C": o["rasa_kolumna_C"],
                   "folder": f"{o['rasa_kolumna_C']} {o['ean']}", "metoda": "rozstrzygnięcie właścicielki", "do_decyzji": "", "rasa_id": rid,
                   "decyzja": "OPTYMALIZUJ", "ean": o["ean"], "klucz_odmiany": o["klucz_odmiany"], "tytul": byid.get(lid, "")}
            if lid not in byid:
                why.append("oferta nie jest aktywna w sklepie")
            if is_done(lid):
                why.append("już zoptymalizowana")
            rr = R[(g0, rid)]
            if any(not rr.get(c) for c in ["PL nazwa"] + [T.L[l]["col"] for l in LANGS if l != "pl"]):
                why.append("brak tłumaczenia nazwy rasy")
            cands.append({"row": row, "lid": lid, "sales": sold.get(lid, 0), "why": why})
    return cands, R


def interleave(sorted_cands):
    """Kolejność: sprzedaż rosnąco; przy równej sprzedaży gatunki na zmianę (pies, kot, koń), żeby partia sprawdzała wszystkie szablony."""
    out, i = [], 0
    while i < len(sorted_cands):
        j = i
        while j < len(sorted_cands) and sorted_cands[j]["sales"] == sorted_cands[i]["sales"]:
            j += 1
        groups = {s: [c for c in sorted_cands[i:j] if c["row"]["gatunek"] == s] for s in ("pies", "kot", "koń")}
        while any(groups.values()):
            for s in groups:
                if groups[s]:
                    out.append(groups[s].pop(0))
        i = j
    return out


def started(n):
    return bool(glob.glob(str(ROOT / "logs" / f"partia_{n:03d}_*.jsonl")))


def load_frozen(n):
    """Zatwierdzona lista partii z plans/partia_00N.csv i _dryrun.json (apply nigdy nie przelicza partii od nowa)."""
    csvp, jp = ROOT / "plans" / f"partia_{n:03d}.csv", ROOT / "plans" / f"partia_{n:03d}_dryrun.json"
    if not csvp.exists() or not jp.exists():
        sys.exit(f"BŁĄD: brak zatwierdzonego planu partii {n} (uruchom 'plan {n}' i zatwierdź)")
    rows = {r["listing_id"]: r for r in csv.DictReader(open(csvp, encoding="utf-8"))}
    batch = []
    for x in json.load(open(jp, encoding="utf-8")):
        r = rows[str(x["listing_id"])]
        batch.append({"lid": x["listing_id"], "en": x["en"], "trans": x["tlumaczenia"],
                      "row": {"ean": x["ean"], "tytul": r["tytul_przed"], "gatunek": x["gatunek"], "rasa_kolumna_C": x["rasa"]}})
    if len(batch) != len(rows):
        sys.exit("BŁĄD: lista CSV i JSON partii różnią się")
    return batch


def plan(n):
    if started(n):
        sys.exit(f"BŁĄD: partia {n} jest już rozpoczęta (są logi), nie nadpisuję zatwierdzonego planu")
    cands, R = build_candidates()
    appr = approved_photo_matches()
    srt = lambda lst: interleave(sorted(lst, key=lambda c: (c["sales"], c["lid"])))
    free = [c for c in cands if not c["why"]]
    # najpierw oferty z zatwierdzonych dopasowań zdjęć (właścicielka kazała je dołączyć do kolejnej partii), potem reszta
    ok = srt([c for c in free if c["row"]["zrodlo"] in appr]) + srt([c for c in free if c["row"]["zrodlo"] not in appr])
    batch, skipped = [], []
    for c in ok:
        if len(batch) == BATCH_MAX:
            break
        en, trans = render(c["row"], c["lid"], R)
        prob = check(en, trans, c["row"]["ean"])
        if not prob:
            miss = reachable(c["row"]["ean"])
            if miss:
                prob.append("zdjęcie niedostępne " + miss)
        if prob:
            c["why"] = prob
            skipped.append(c)
            continue
        c.update(en=en, trans=trans)
        batch.append(c)
    with open(STATUS_CSV, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["zrodlo", "listing_id", "rasa", "gatunek", "sprzedaz_szt", "status", "powod"])
        inb = {c["lid"] for c in batch}
        for c in cands:
            r = c["row"]
            st = f"partia {n}" if c["lid"] in inb else ("kandydat" if not c["why"] else "wykluczona")
            w.writerow([r["zrodlo"], c["lid"], r["rasa_kolumna_C"], r["gatunek"], c["sales"], st, "; ".join(c["why"])])
    p = ROOT / "plans" / f"partia_{n:03d}.csv"
    with open(p, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["lp", "listing_id", "zrodlo", "gatunek", "rasa", "ean", "sprzedaz_szt", "tytul_przed", "tytul_po", "wzorzec_tytulu", "ABC", "zdrobnienie", "synonim", "odmiana", "tagi"])
        for i, c in enumerate(batch, 1):
            r, e = c["row"], c["en"]
            w.writerow([i, c["lid"], r["zrodlo"], r["gatunek"], r["rasa_kolumna_C"], r["ean"], c["sales"], r["tytul"], e["title"],
                        e["wzorzec_tytulu"], "".join(f"{k}{v + 1}" for k, v in e["ABC"].items()), e["nick"], e["synonim"], e["odmiana"], "; ".join(e["tags"])])
    dry = [{"listing_id": c["lid"], "zrodlo": c["row"]["zrodlo"], "ean": c["row"]["ean"], "gatunek": c["row"]["gatunek"],
            "rasa": c["row"]["rasa_kolumna_C"], "en": c["en"], "tlumaczenia": c["trans"]} for c in batch]
    (ROOT / "plans" / f"partia_{n:03d}_dryrun.json").write_text(json.dumps(dry, ensure_ascii=False, indent=1))
    print(f"partia {n}: {len(batch)} ofert; kandydatów łącznie {len(ok)}; wykluczonych w tym kroku {len(skipped)}")
    print("sprzedaż w partii (szt.): min", min(c["sales"] for c in batch), "max", max(c["sales"] for c in batch))
    from collections import Counter
    print("gatunki:", dict(Counter(c["row"]["gatunek"] for c in batch)))
    print("wykluczone ogółem (powody):", dict(Counter(w.split(":")[0] for c in cands for w in c["why"])))
    return batch, R


# ---------- zapis ----------
def steps_for(c):
    en, trans, lid = c["en"], c["trans"], c["lid"]
    b = f"/v3/application/shops/{SHOP_ID}/listings/{lid}"
    s = [{"id": "listing", "method": "PATCH", "path": b, "kind": "form",
          "fields": {"title": en["title"], "description": en["description"], "tags": ",".join(en["tags"]), "materials": "ceramic"}},
         {"id": "wlasciwosc_materials", "method": "PUT", "path": f"{b}/properties/{MATERIAL_PROPERTY}", "kind": "form",
          "fields": {"value_ids": str(CERAMIC_VALUE), "values": "Ceramic"}},
         {"id": "how_its_made", "method": "PATCH", "path": b, "kind": "form",
          "fields": {"who_made": "collective", "when_made": "made_to_order", "is_supply": "false"}},
         {"id": "personalizacja", "method": "POST", "path": b + "/personalization", "kind": "json",
          "fields": {"personalization_questions": base.PERSONALIZATION}}]
    for rank, f in enumerate(IMAGE_ORDER, 1):
        s.append({"id": f"zdjecie_{rank}_{f}", "method": "POST", "path": b + "/images", "kind": "multipart",
                  "fields": {"rank": str(rank), "overwrite": "true" if rank == 1 else "false", "alt_text": en["alts"][f]},
                  "file": ("image", f)})
    s.append({"id": "wideo", "method": "POST", "path": b + "/videos", "kind": "multipart",
              "fields": {"name": base.VIDEO.name}, "file": ("video", str(base.VIDEO))})
    for code in LANGS:
        d = trans[code]
        s.append({"id": f"tlumaczenie_{code}", "method": "POST", "path": b + f"/translations/{code}", "kind": "form",
                  "fields": {"title": d["title"], "description": d["description"], "tags": ",".join(d["tags"])}})
    return s


def fetch_batch_image(ean, f):
    if f == "original.jpg":
        return ROOT / "media" / "plytki_with_photo" / "original.jpg"
    p = CACHE / ean / f
    if not p.exists():
        p.parent.mkdir(parents=True, exist_ok=True)
        with urllib.request.urlopen(urllib.request.Request(f"{EAN_BASE}{ean}/{f}", headers=UA), timeout=120) as r:
            p.write_bytes(r.read())
    return p


def send(step, ean):
    if step["kind"] == "multipart" and step["file"][0] == "image":
        part, src = step["file"]
        path = fetch_batch_image(ean, src)
        data, ct = base.multipart(step["fields"], (part, path.name), path.read_bytes(), "image/jpeg")
        return api.request(step["method"], step["path"], headers={"Content-Type": ct}, data=data, retries=1)
    return base.send(step)


def zlog(lid, step_id, before, after, status):
    f = ROOT / "logs" / f"zapis_{lid}_{TODAY}.jsonl"
    with open(f, "a") as fh:
        fh.write(json.dumps({"listing_id": lid, "id": step_id, "field": step_id, "before": before, "after": after, "status": status,
                             "time": datetime.datetime.now().isoformat(timespec="seconds")}, ensure_ascii=False) + "\n")


def done_steps(lid):
    d = set()
    for f in glob.glob(str(ROOT / "logs" / f"zapis_{lid}_*.jsonl")):
        d |= {j["id"] for j in map(json.loads, open(f).read().splitlines()) if j["status"] == "ok"}
    return d


def backup_unique(lid):
    live = api.get(f"/v3/application/listings/{lid}?includes=Images,Videos,Translations,Personalization")
    out = {"listing": live,
           "properties": api.get(f"/v3/application/shops/{SHOP_ID}/listings/{lid}/properties"),
           "inventory": api.get(f"/v3/application/listings/{lid}/inventory")}
    p = ROOT / "backup" / f"{lid}_{TODAY}.json"
    k = 1
    while p.exists():
        k += 1
        p = ROOT / "backup" / f"{lid}_{TODAY}_{k}.json"
    p.write_text(json.dumps(out, ensure_ascii=False, indent=1))
    return p, live


def preflight(c, live):
    """Oferta ma być w stanie z eksportu: ten sam tytuł, 1 zdjęcie, brak wideo i personalizacji."""
    pr = []
    if "error" in live:
        return ["odczyt oferty: " + str(live["error"])[:100]]
    if html.unescape(live["title"]) != c["row"]["tytul"]:
        pr.append("tytuł w Etsy inny niż w eksporcie")
    if len(live.get("images", [])) != 1:
        pr.append(f"zdjęć: {len(live.get('images', []))} (oczekiwano 1)")
    if live.get("videos"):
        pr.append("oferta ma już wideo")
    if live.get("is_personalizable"):
        pr.append("oferta ma już personalizację")
    if live["state"] != "active":
        pr.append("status " + live["state"])
    return pr


def apply_batch(n):
    tok = json.loads(api.TOKEN.read_text())
    if "listings_w" not in tok.get("scope", ""):
        sys.exit("BŁĄD: token bez listings_w")
    batch = load_frozen(n)
    print(f"partia {n}: {len(batch)} ofert z zatwierdzonego planu (bez przeliczania)")
    if len(batch) > BATCH_MAX:
        sys.exit("BŁĄD: partia większa niż 50")
    summary = ROOT / "logs" / f"partia_{n:03d}_{TODAY}.jsonl"
    skip = {int(x) for a in sys.argv if a.startswith("--skip=") for x in a[7:].split(",")}
    for i, c in enumerate(batch, 1):
        lid = c["lid"]
        if lid in skip:
            print(i, lid, "POMINIĘTA na żądanie (--skip), do poprawki")
            continue
        if is_done(lid):
            print(i, lid, "pomijam (już zrobiona)")
            continue
        p, live = backup_unique(lid)
        pr = preflight(c, live)
        if pr and not (done_steps(lid) - {"weryfikacja"}):
            print(i, lid, "ZATRZYMANA przed zapisem:", "; ".join(pr))
            open(summary, "a").write(json.dumps({"listing_id": lid, "status": "pominięta", "powod": pr}, ensure_ascii=False) + "\n")
            continue
        done = done_steps(lid)
        resumed = bool(done - {"weryfikacja"})
        for s in steps_for(c):
            if s["id"] in done:
                continue
            if resumed and s["id"].startswith("zdjecie_"):  # wznowienie: zdjęcie mogło już wejść, a nie zdążyć się zalogować (bez duplikatów)
                imgs = api.get(f"/v3/application/listings/{lid}?includes=Images").get("images", [])
                if any(i["rank"] == int(s["fields"]["rank"]) and (i.get("alt_text") or "") == s["fields"]["alt_text"] for i in imgs):
                    zlog(lid, s["id"], "backup/" + p.name, "już w ofercie (wykryte przy wznowieniu)", "ok")
                    print(f"  {lid} {s['id']}: już w ofercie, pomijam")
                    continue
            resp = send(s, c["row"]["ean"])
            if s["id"].startswith("tlumaczenie_") and not base.ok(resp) and "exist" in str(resp).lower():
                resp = send({**s, "method": "PUT"}, c["row"]["ean"])
            st = "ok" if base.ok(resp) else "błąd"
            zlog(lid, s["id"], "backup/" + p.name, {k: (v if len(str(v)) < 200 else str(v)[:200] + "...") for k, v in s["fields"].items()}, st)
            if st != "ok":
                open(summary, "a").write(json.dumps({"listing_id": lid, "status": "błąd", "krok": s["id"], "odpowiedź": str(resp)[:300]}, ensure_ascii=False) + "\n")
                sys.exit(f"{lid}: BŁĄD w kroku {s['id']}: {str(resp)[:300]}\nPrzerywam całą partię; po wyjaśnieniu uruchom ponownie (wykonane kroki zostaną pominięte).")
            time.sleep(0.15)
        for attempt in range(4):  # odczyty Etsy bywają niespójne tuż po zapisie: powtarzamy, zanim uznamy niezgodność
            after = api.get(f"/v3/application/listings/{lid}?includes=Images,Videos,Translations,Personalization")
            chk = {"cena_ok": after["price"] == live["price"], "stan_ok": after["quantity"] == live["quantity"], "status_ok": after["state"] == live["state"],
                   "zdjec": len(after.get("images", [])), "wideo": len(after.get("videos", [])),
                   "pytan": len(after["personalization"]["personalization_questions"]),
                   "tlumaczen": sum(1 for k, v in after["translations"].items() if v and k != "en-US"),
                   "who_made": after["who_made"], "proba": attempt + 1}
            good = chk["cena_ok"] and chk["stan_ok"] and chk["status_ok"] and chk["zdjec"] == 7 and chk["wideo"] == 1 and chk["pytan"] == 5 \
                and chk["tlumaczen"] == 8 and chk["who_made"] == "collective"
            if good:
                break
            time.sleep(3)
        zlog(lid, "weryfikacja", {"cena": live["price"], "stan": live["quantity"]}, chk, "ok" if good else "niezgodność")
        open(summary, "a").write(json.dumps({"listing_id": lid, "status": "ok" if good else "niezgodność", **chk}, ensure_ascii=False) + "\n")
        print(i, lid, "ok" if good else "NIEZGODNOŚĆ " + str(chk))
        if not good:
            sys.exit("Weryfikacja niezgodna, przerywam partię.")
    print("Partia zakończona. Proszę sprawdzić próbkę w Etsy przed następną partią.")


def main():
    cmd = sys.argv[1]
    if cmd == "ids":
        fetch_ids()
    elif cmd == "plan":
        plan(int(sys.argv[2]))
    elif cmd == "apply":
        if "--apply" not in sys.argv:
            sys.exit("Brak flagi --apply: nic nie zapisuję. Użyj 'plan N' do dry-run.")
        apply_batch(int(sys.argv[2]))


if __name__ == "__main__":
    main()
